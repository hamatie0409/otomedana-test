# -*- coding: utf-8 -*-
"""DBの全テーブルをExcelブック1つに書き出す。

    python3 scripts/export_db_xlsx.py
    python3 scripts/export_db_xlsx.py --out /path/to/file.xlsx

DBは60MBのSQLiteで、SQLを書かないと中身が見えない。中身を人が眺めたり
他のツールに渡したりできるよう、テーブルごとに1シートで丸ごと出す。

先頭に「目次」シートを置き、どのテーブルが何行あるかと、主な列の意味を書く。
"""
import argparse
import os
import re
import sqlite3
import sys
import datetime

from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(ROOT, "data", "vndb_otome.db")
OUT = os.path.join(ROOT, "data", "otomedana_db.xlsx")

CELL_MAX = 32000          # Excelの上限は32,767文字。余裕をみて切る
BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")   # Excelが受け付けない制御文字

# 目次に出す説明。書いていないテーブルは空欄で出る
NOTE = {
    "games": "作品。サイトの1ページに対応するのは slugs が kind='game' のもの",
    "editions": "版（機種×通常版/限定版/DL版）。購入リンクはここに付く",
    "characters": "キャラクター。声優・属性つき",
    "traits": "キャラクターの属性。髪・瞳・性格など。category と組で意味が決まる",
    "vn_tags": "作品のタグ。ネタバレ度の高いものは除外済み",
    "staff_credits": "シナリオ・原画・音楽などのスタッフ",
    "offers": "購入先。channel（楽天/Amazon/駿河屋/アニメイト/メルカリ）ごとに1行",
    "slugs": "URLスラッグ。is_page=1 が実際に生成されるページ",
    "shop_images": "作品の表紙に採用した商品写真。width/height は実測値",
    "image_probe": "商品写真の実寸の測定結果（キャッシュ）",
    "rakuten_items": "楽天でJAN検索して取れた商品",
    "rakuten_title_items": "JANの無い版をタイトル検索で補ったもの",
    "producers": "発売元の分類。lang と type で有志翻訳グループを見分ける",
    "corrections_log": "corrections/works.csv から適用した訂正の記録。出典つき",
    "suruga_dates": "駿河屋のJAN照会で得た発売日（キャッシュ）",
    "vndb_links": "外部リンク。公式サイト・ストア・ErogameScapeなど",
    "relations": "関連作品",
    "languages": "対応言語",
    "platforms": "対応機種",
    "gtins": "JANコード",
    "shop_urls": "店ごとの検索URL",
}


def clean(v):
    if v is None or isinstance(v, (int, float)):
        return v
    s = str(v)
    s = BAD.sub("", s)
    return s[:CELL_MAX] + "…" if len(s) > CELL_MAX else s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=OUT)
    a = ap.parse_args()
    if not os.path.exists(DB):
        sys.exit("DBがありません: %s" % DB)

    con = sqlite3.connect(DB)
    tables = [r[0] for r in con.execute(
        "SELECT name FROM sqlite_master WHERE type='table' "
        "AND name NOT LIKE 'sqlite_%' ORDER BY name")]

    # 大きいので write_only。行を書いたそばから捨てるためメモリが増えない
    wb = Workbook(write_only=True)
    idx = wb.create_sheet("目次")
    idx.append(["オトメ棚 データベース書き出し",
                datetime.datetime.now().strftime("%Y-%m-%d %H:%M")])
    idx.append([])
    idx.append(["テーブル", "行数", "列数", "説明"])

    total = 0
    rows_meta = []
    for t in tables:
        cols = [d[1] for d in con.execute("PRAGMA table_info(%s)" % t)]
        n = con.execute("SELECT COUNT(*) FROM %s" % t).fetchone()[0]
        rows_meta.append((t, n, len(cols), NOTE.get(t, "")))
        total += n
    for m in rows_meta:
        idx.append(list(m))
    idx.append([])
    idx.append(["合計", total, "", "%d テーブル" % len(tables)])

    for t in tables:
        cols = [d[1] for d in con.execute("PRAGMA table_info(%s)" % t)]
        ws = wb.create_sheet(t[:31])
        ws.append(cols)
        n = 0
        for r in con.execute("SELECT * FROM %s" % t):
            ws.append([clean(v) for v in r])
            n += 1
        print("  %-22s %8d行" % (t, n))

    wb.save(a.out)
    mb = os.path.getsize(a.out) / 1048576.0
    print()
    print("  書き出し: %s  (%.1f MB / %d行)" % (os.path.relpath(a.out, ROOT), mb, total))


if __name__ == "__main__":
    main()
